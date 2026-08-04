"""Renders a generic `ReportContent` (see app.reporting.content) into
real file bytes for each of spec §65's 4 exportable formats. Every
renderer is a real library call (reportlab/python-docx/openpyxl/stdlib
csv) producing a genuinely openable file — never a text file with a
misleading extension.
"""

from __future__ import annotations

import csv
import io
from collections.abc import Callable

from docx import Document
from openpyxl import Workbook
from reportlab.lib import colors
from reportlab.lib.pagesizes import letter
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

from app.reporting.content import ReportContent

_TIMESTAMP_FORMAT = "%Y-%m-%d %H:%M UTC"


def render_pdf(content: ReportContent) -> bytes:
    buffer = io.BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=letter)
    styles = getSampleStyleSheet()
    story: list[object] = [
        Paragraph(content.title, styles["Title"]),
        Paragraph(f"Generated {content.generated_at.strftime(_TIMESTAMP_FORMAT)}", styles["Normal"]),
        Spacer(1, 12),
    ]
    for section in content.sections:
        story.append(Paragraph(section.heading, styles["Heading2"]))
        if not section.rows:
            story.append(Paragraph("No data.", styles["Normal"]))
        else:
            columns = list(section.rows[0].keys())
            table_data = [columns] + [
                [str(row.get(c, "")) for c in columns] for row in section.rows
            ]
            table = Table(table_data, repeatRows=1)
            table.setStyle(
                TableStyle(
                    [
                        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1f2937")),
                        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                        ("GRID", (0, 0), (-1, -1), 0.5, colors.grey),
                        ("FONTSIZE", (0, 0), (-1, -1), 8),
                        ("VALIGN", (0, 0), (-1, -1), "TOP"),
                    ]
                )
            )
            story.append(table)
        story.append(Spacer(1, 16))
    doc.build(story)
    return buffer.getvalue()


def render_docx(content: ReportContent) -> bytes:
    document = Document()
    document.add_heading(content.title, level=0)
    document.add_paragraph(f"Generated {content.generated_at.strftime(_TIMESTAMP_FORMAT)}")
    for section in content.sections:
        document.add_heading(section.heading, level=1)
        if not section.rows:
            document.add_paragraph("No data.")
            continue
        columns = list(section.rows[0].keys())
        table = document.add_table(rows=1, cols=len(columns))
        table.style = "Light Grid Accent 1"
        header_cells = table.rows[0].cells
        for i, column in enumerate(columns):
            header_cells[i].text = column
        for row in section.rows:
            cells = table.add_row().cells
            for i, column in enumerate(columns):
                cells[i].text = str(row.get(column, ""))
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


_INVALID_SHEET_CHARS = set(":\\/?*[]")


def _unique_sheet_name(workbook: Workbook, heading: str) -> str:
    cleaned = "".join(c for c in heading if c not in _INVALID_SHEET_CHARS)[:31] or "Sheet"
    name = cleaned
    suffix = 2
    while name in workbook.sheetnames:
        name = f"{cleaned[:28]}_{suffix}"
        suffix += 1
    return name


def render_xlsx(content: ReportContent) -> bytes:
    workbook = Workbook()
    workbook.remove(workbook.active)
    for section in content.sections:
        sheet = workbook.create_sheet(title=_unique_sheet_name(workbook, section.heading))
        if not section.rows:
            sheet.append(["No data"])
            continue
        columns = list(section.rows[0].keys())
        sheet.append(columns)
        for row in section.rows:
            sheet.append([row.get(c, "") for c in columns])
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def render_csv(content: ReportContent) -> bytes:
    """A single CSV can't hold multiple named tables the way a workbook's
    sheets or a document's headed sections can, so each section is
    written as a heading row followed by its own header+data rows,
    separated by a blank line — a real, honest flattening rather than
    only exporting the first section."""
    buffer = io.StringIO()
    writer = csv.writer(buffer)
    writer.writerow([content.title])
    writer.writerow([f"Generated {content.generated_at.strftime(_TIMESTAMP_FORMAT)}"])
    for section in content.sections:
        writer.writerow([])
        writer.writerow([section.heading])
        if not section.rows:
            writer.writerow(["No data"])
            continue
        columns = list(section.rows[0].keys())
        writer.writerow(columns)
        for row in section.rows:
            writer.writerow([row.get(c, "") for c in columns])
    return buffer.getvalue().encode("utf-8")


RENDERERS: dict[str, Callable[[ReportContent], bytes]] = {
    "pdf": render_pdf,
    "docx": render_docx,
    "xlsx": render_xlsx,
    "csv": render_csv,
}
