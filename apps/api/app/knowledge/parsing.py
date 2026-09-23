"""Text extraction per content type (spec section 46, Document Intelligence).

Only formats with a genuine extractor are attempted. Anything else raises
UnsupportedFormatError so the caller can mark the document
`unsupported_format` rather than silently producing empty/fake content
(spec section 82, No Fake Features).
"""

from __future__ import annotations

import csv
import io
from dataclasses import dataclass

TEXT_CONTENT_TYPES = {"text/plain", "text/markdown", "text/csv"}
PDF_CONTENT_TYPES = {"application/pdf"}
DOCX_CONTENT_TYPES = {"application/vnd.openxmlformats-officedocument.wordprocessingml.document"}
XLSX_CONTENT_TYPES = {"application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"}
HTML_CONTENT_TYPES = {"text/html"}


class UnsupportedFormatError(Exception):
    def __init__(self, content_type: str) -> None:
        super().__init__(f"No text extractor available for content type '{content_type}'")
        self.content_type = content_type


@dataclass(frozen=True)
class ExtractedDocument:
    """Extracted text, plus where its pages begin.

    `page_starts` holds the character offset at which each page starts, so
    a citation can say "page 4" instead of a chunk number. It is empty for
    formats that genuinely have no pages — a Markdown file or a .docx
    paragraph stream does not, and a made-up page number sends the reader
    somewhere that does not exist.
    """

    text: str
    page_starts: tuple[int, ...] = ()


def extract_document(content_type: str, data: bytes) -> ExtractedDocument:
    if content_type in TEXT_CONTENT_TYPES:
        return ExtractedDocument(_extract_text_plain(content_type, data))
    if content_type in PDF_CONTENT_TYPES:
        return _extract_pdf(data)
    if content_type in DOCX_CONTENT_TYPES:
        return ExtractedDocument(_extract_docx(data))
    if content_type in XLSX_CONTENT_TYPES:
        return ExtractedDocument(_extract_xlsx(data))
    if content_type in HTML_CONTENT_TYPES:
        return ExtractedDocument(_extract_html(data))
    raise UnsupportedFormatError(content_type)


def extract_text(content_type: str, data: bytes) -> str:
    """Just the text, for the callers that only ever wanted that."""
    return extract_document(content_type, data).text


def _extract_text_plain(content_type: str, data: bytes) -> str:
    text = data.decode("utf-8", errors="replace")
    if content_type == "text/csv":
        # Normalize CSV into readable row text rather than raw comma soup,
        # so chunking/embedding operates on something coherent.
        reader = csv.reader(io.StringIO(text))
        return "\n".join(", ".join(row) for row in reader)
    return text


PAGE_SEPARATOR = "\n\n"


def _extract_pdf(data: bytes) -> ExtractedDocument:
    from pypdf import PdfReader

    reader = PdfReader(io.BytesIO(data))
    pages = [page.extract_text() or "" for page in reader.pages]

    # Assembled by hand rather than with join() because the offsets are
    # the point: a citation resolves its page by asking which of these the
    # passage starts after.
    starts: list[int] = []
    parts: list[str] = []
    offset = 0
    for index, page in enumerate(pages):
        if index:
            parts.append(PAGE_SEPARATOR)
            offset += len(PAGE_SEPARATOR)
        starts.append(offset)
        parts.append(page)
        offset += len(page)
    return ExtractedDocument("".join(parts), tuple(starts))


def _extract_docx(data: bytes) -> str:
    """Everything readable in the document, tables included.

    `document.paragraphs` is only the top-level paragraph stream: it
    excludes every cell of every table. That is not a small omission in
    compliance documents, where the facts are routinely *in* the tables
    — a RoPA's activity name, purpose, data subjects and retention; a
    DPIA's risk register; a transfer's destination and mechanism. A
    document whose tables were dropped reached evidence analysis, the
    Copilot index and privacy extraction with its headings intact and
    its content missing, which reads as an empty section rather than as
    a failure.

    Body children are walked in document order so a table stays with the
    heading that introduces it, and each row is joined into one line so
    a two-column "label | value" table reads back as "label: value" —
    which is what it means, and what any reader downstream expects.
    """
    import docx
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    document = docx.Document(io.BytesIO(data))
    lines: list[str] = []

    for child in document.element.body.iterchildren():
        tag = child.tag.split("}")[-1]
        if tag == "p":
            text = Paragraph(child, document).text.strip()
            if text:
                lines.append(text)
        elif tag == "tbl":
            for row in Table(child, document).rows:
                cells = [cell.text.strip() for cell in row.cells]
                # A repeated cell is how python-docx reports a merge; one
                # value said twice is noise in the text stream.
                deduped: list[str] = []
                for cell in cells:
                    if cell and (not deduped or deduped[-1] != cell):
                        deduped.append(cell)
                if not deduped:
                    continue
                # Two columns are a label and its value almost every
                # time, so they are written the way the rest of the
                # pipeline already reads: "Purpose: …".
                if len(deduped) == 2:
                    lines.append(f"{deduped[0]}: {deduped[1]}")
                else:
                    lines.append(" | ".join(deduped))

    return "\n".join(lines)


def _extract_xlsx(data: bytes) -> str:
    from openpyxl import load_workbook

    workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    lines: list[str] = []
    for sheet in workbook.worksheets:
        lines.append(f"# {sheet.title}")
        for row in sheet.iter_rows(values_only=True):
            cells = [str(cell) for cell in row if cell is not None]
            if cells:
                lines.append(", ".join(cells))
    return "\n".join(lines)


def _extract_html(data: bytes) -> str:
    from bs4 import BeautifulSoup

    soup = BeautifulSoup(data.decode("utf-8", errors="replace"), "html.parser")
    for tag in soup(["script", "style"]):
        tag.decompose()
    return str(soup.get_text(separator="\n", strip=True))
