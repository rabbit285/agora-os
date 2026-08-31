"""What a customer receives, and what they must never receive.

The feature is not "generate a compliance PDF". It is: an approved
review, a choice of customer-facing output, an inspection of the exact
contents, a chance to remove what should not go, and only then an
immutable deliverable. These pin every step of that, and the two things
that make it safe rather than merely convenient.

**Preview writes nothing.** Opening and refreshing must not leave a trail
of packages nobody sent, or the export history stops meaning "what we
sent" and starts meaning "what somebody looked at".

**Nothing leaves that nobody cleared.** A document's disclosure is the
gate. Unclassified is not shareable, restricted is refused outright, and
another customer's information never appears at all.
"""

from __future__ import annotations

import uuid
import zipfile
from io import BytesIO

import pytest
from httpx import AsyncClient
from openpyxl import load_workbook
from pypdf import PdfReader
from sqlalchemy import func, select
from sqlalchemy.ext.asyncio import AsyncSession

from app.audit.models import AuditEvent
from app.compliance import review_export
from app.compliance.models import (
    Evidence,
    Policy,
    Questionnaire,
    QuestionnaireAnswer,
    QuestionnaireQuestion,
    ReviewExport,
    SecurityReview,
)
from app.compliance.review_render import (
    render_evidence_index,
    render_questionnaire_xlsx,
    render_summary_pdf,
)
from app.core.models import File


async def _signup(
    client: AsyncClient, email: str, org: str
) -> tuple[uuid.UUID, uuid.UUID]:
    """Returns (organization_id, user_id).

    The user matters: `File.uploaded_by` is not nullable, and a document
    with no uploader is not a shape this product ever produces.
    """
    response = await client.post(
        "/api/v1/auth/signup",
        json={
            "email": email,
            "password": "correct-horse-battery",
            "full_name": "Zach Warren",
            "organization_name": org,
        },
    )
    assert response.status_code == 200, response.text
    body = response.json()
    return (
        uuid.UUID(body["organization"]["id"]),
        uuid.UUID(body["user"]["id"]),
    )


ANSWERS = [
    (
        "Do you have an incident-response process?",
        "Yes. The organization maintains a documented incident-response process; "
        "however, current evidence does not demonstrate completion of the most "
        "recent tabletop exercise.",
    ),
    (
        "Do you encrypt customer data at rest?",
        "Yes. All customer data is encrypted at rest using AES-256 with managed keys.",
    ),
    (
        "Do you maintain an information security policy?",
        "Yes. The Information Security Policy is reviewed annually and approved "
        "by leadership.",
    ),
]


async def _review(
    client: AsyncClient,
    db: AsyncSession,
    *,
    customer: str = "Agora Origin",
    approve: bool = True,
) -> SecurityReview:
    """A finished customer review, built the way the product builds one."""
    created = await client.post(
        "/api/v1/compliance/questionnaires",
        json={
            "title": f"{customer} Security Questionnaire",
            "source": "customer_provided",
            "customer_name": customer,
        },
    )
    assert created.status_code in (200, 201), created.text
    questionnaire_id = created.json()["id"]

    for text, answer in ANSWERS:
        question = await client.post(
            f"/api/v1/compliance/questionnaires/{questionnaire_id}/questions",
            json={"text": text, "category": "security"},
        )
        await client.patch(
            f"/api/v1/compliance/questionnaires/{questionnaire_id}/questions/"
            f"{question.json()['id']}/answer",
            json={
                "answer_text": answer,
                "status": "approved" if approve else "draft",
            },
        )

    review = (
        await db.execute(
            select(SecurityReview).where(SecurityReview.customer_name == customer)
        )
    ).scalars().first()
    assert review is not None
    await db.refresh(review)
    return review


async def _document(
    db: AsyncSession,
    org_id: uuid.UUID,
    user_id: uuid.UUID,
    *,
    title: str,
    disclosure: str,
    as_policy: bool = True,
    policy_status: str = "approved",
    version: str = "2.0",
) -> File:
    file_record = File(
        organization_id=org_id,
        storage_key=f"org/{org_id}/{uuid.uuid4().hex}.pdf",
        original_filename=f"{title}.pdf",
        content_type="application/pdf",
        size_bytes=1024,
        uploaded_by=user_id,
        disclosure=disclosure,
    )
    db.add(file_record)
    await db.flush()
    if as_policy:
        db.add(
            Policy(
                organization_id=org_id,
                title=title,
                file_id=file_record.id,
                status=policy_status,
                version=version,
            )
        )
    else:
        db.add(
            Evidence(
                organization_id=org_id,
                title=title,
                file_id=file_record.id,
            )
        )
    await db.flush()
    return file_record


# --- §14 qualified answers stay qualified ------------------------------


@pytest.mark.parametrize(
    ("text", "expected"),
    [
        (
            "Yes. We maintain a documented procedure; however, the most recent "
            "tabletop exercise is not evidenced.",
            True,
        ),
        ("Partially implemented across production systems.", True),
        ("Encryption at rest is planned for Q4.", True),
        ("Yes. All data is encrypted at rest with AES-256.", False),
        ("We operate a 24/7 security operations centre.", False),
    ],
)
def test_a_limitation_in_an_answer_is_recognised(text: str, expected: bool) -> None:
    """Report generation must not turn a qualified answer into "Yes".

    That is the failure that makes assurance documents worthless, and it
    is easy to do by accident when summarising.
    """
    assert review_export.is_qualified(text) is expected


# --- §3/§22 preview comes first, and writes nothing --------------------


@pytest.mark.anyio
async def test_preview_shows_the_contents_and_creates_nothing(
    client: AsyncClient, db_session: AsyncSession
) -> None:
    """§28: preview first, and the history stays empty until Generate."""
    _org, _user = await _signup(client, f"prev-{uuid.uuid4().hex[:8]}@example.com", "Prev Org")
    review = await _review(client, db_session)

    for _ in range(3):
        response = await client.post(
            f"/api/v1/compliance/security-reviews/{review.id}/export/preview",
            json={"output_type": "summary"},
        )
        assert response.status_code == 200, response.text

    preview = response.json()
    assert preview["output_label"] == "Security Review Summary"
    assert preview["file_format"] == "pdf"
    assert preview["customer_name"] == "Agora Origin"
    assert preview["progress"]["approved"] == 3
    assert preview["progress"]["percent_complete"] == 100
    assert len(preview["answers"]) == 3
    # §14: the qualified answer is surfaced as a disclosure.
    assert len(preview["disclosures"]) == 1
    assert "tabletop" in preview["disclosures"][0]
    assert preview["ready"] is True
    assert preview["blockers"] == []
    assert preview["disclaimer"].startswith("This report summarizes")
    assert preview["next_version"] == 1
    # The fingerprint the client sends back to confirm.
    assert preview["digest"]

    # §22/§29: three previews, no records, no business events.
    exports = int(
        (
            await db_session.execute(
                select(func.count(ReviewExport.id)).where(
                    ReviewExport.security_review_id == review.id
                )
            )
        ).scalar_one()
    )
    assert exports == 0
    events = list(
        (
            await db_session.execute(
                select(AuditEvent).where(
                    AuditEvent.action.like("compliance.review_export.%")
                )
            )
        )
        .scalars()
        .all()
    )
    assert events == []


@pytest.mark.anyio
async def test_the_preview_never_shows_internal_numbers(
    client: AsyncClient, db_session: AsyncSession
) -> None:
    """§7/§8: no confidence, no chunk ids, no model reasoning.

    Confidence is a number about our process, not a fact about the
    customer's security, and a customer reading "0.82" learns something
    misleading.
    """
    _org, _user = await _signup(client, f"clean-{uuid.uuid4().hex[:8]}@example.com", "Clean Org")
    review = await _review(client, db_session)

    preview = (
        await client.post(
            f"/api/v1/compliance/security-reviews/{review.id}/export/preview",
            json={"output_type": "summary"},
        )
    ).json()

    for answer in preview["answers"]:
        assert "confidence" not in answer
        assert "citations" not in answer
        assert "chunk_id" not in answer
        assert "review_reason" not in answer


# --- §6 approved answers only ------------------------------------------


@pytest.mark.anyio
async def test_an_unapproved_answer_blocks_the_export(
    client: AsyncClient, db_session: AsyncSession
) -> None:
    _org, _user = await _signup(client, f"unap-{uuid.uuid4().hex[:8]}@example.com", "Unap Org")
    review = await _review(client, db_session, approve=False)

    preview = (
        await client.post(
            f"/api/v1/compliance/security-reviews/{review.id}/export/preview",
            json={"output_type": "summary"},
        )
    ).json()

    assert preview["ready"] is False
    codes = {blocker["code"] for blocker in preview["blockers"]}
    assert "unanswered" in codes
    # And nothing unapproved is quoted.
    assert preview["answers"] == []


@pytest.mark.anyio
async def test_an_answer_edited_after_approval_blocks_the_export(
    client: AsyncClient, db_session: AsyncSession
) -> None:
    """§6: neither the old state nor the new one is safe to send.

    The check compares the current text against the text that was signed
    off. Timestamps cannot answer it — `updated_at` moves on the same
    flush that records the approval, which flagged three freshly approved
    answers as changed the first time this was built.
    """
    _org, _user = await _signup(client, f"edit-{uuid.uuid4().hex[:8]}@example.com", "Edit Org")
    review = await _review(client, db_session)

    fresh = (
        await client.post(
            f"/api/v1/compliance/security-reviews/{review.id}/export/preview",
            json={"output_type": "summary"},
        )
    ).json()
    assert fresh["ready"] is True, "a freshly approved review is not 'changed'"

    answer = (
        await db_session.execute(select(QuestionnaireAnswer).limit(1))
    ).scalars().first()
    answer.answer_text = "Reworded without going back through approval."
    await db_session.commit()

    preview = (
        await client.post(
            f"/api/v1/compliance/security-reviews/{review.id}/export/preview",
            json={"output_type": "summary"},
        )
    ).json()
    codes = {blocker["code"] for blocker in preview["blockers"]}
    assert "changed_since_approval" in codes
    assert preview["ready"] is False


# --- §8/§10/§25 nothing leaves that nobody cleared ----------------------


@pytest.mark.anyio
async def test_only_cleared_documents_are_proposed(
    client: AsyncClient, db_session: AsyncSession
) -> None:
    """§25, exactly: four documents, four outcomes.

    The shareable policy is offered. The unclassified architecture
    diagram is not — an undecided document is not a shareable one. The
    restricted vulnerability report is refused outright. And another
    customer's file never appears, because it belongs to another
    organization and this query cannot see it.
    """
    org_id, user_id = await _signup(client, f"leak-{uuid.uuid4().hex[:8]}@example.com", "Leak Org")

    await _document(
        db_session, org_id, user_id, title="Information Security Policy",
        disclosure="customer_shareable",
    )
    await _document(
        db_session, org_id, user_id, title="Internal architecture diagram",
        disclosure="internal", as_policy=False,
    )
    await _document(
        db_session, org_id, user_id, title="Restricted vulnerability assessment",
        disclosure="restricted", as_policy=False,
    )
    # Another tenant's document, marked shareable in *their* org.
    other_org = uuid.uuid4()
    review = await _review(client, db_session)
    db_session.add(
        File(
            organization_id=other_org,
            storage_key=f"org/{other_org}/{uuid.uuid4().hex}.pdf",
            original_filename="Customer B security report.pdf",
            content_type="application/pdf",
            size_bytes=10,
            uploaded_by=user_id,
            # Marked shareable in the other tenant. It still must never
            # appear here — the org boundary is what keeps it out, not
            # its disclosure state.
            disclosure="customer_shareable",
        )
    )
    await db_session.commit()

    preview = (
        await client.post(
            f"/api/v1/compliance/security-reviews/{review.id}/export/preview",
            json={"output_type": "evidence_pack"},
        )
    ).json()

    by_title = {document["title"]: document for document in preview["documents"]}
    assert "Customer B security report" not in " ".join(by_title)

    policy = by_title["Information Security Policy"]
    assert policy["eligible"] is True
    assert policy["included"] is True

    internal = by_title["Internal architecture diagram"]
    assert internal["eligible"] is False
    assert internal["included"] is False
    assert "not currently approved for external sharing" in internal["reason"]

    restricted = by_title["Restricted vulnerability assessment"]
    assert restricted["eligible"] is False
    assert restricted["included"] is False
    assert "cannot be shared externally" in restricted["reason"]


@pytest.mark.anyio
async def test_an_uncleared_document_cannot_be_forced_in(
    client: AsyncClient, db_session: AsyncSession
) -> None:
    """§10: asking for it is not the same as being allowed it."""
    org_id, user_id = await _signup(client, f"force-{uuid.uuid4().hex[:8]}@example.com", "Force Org")
    restricted = await _document(
        db_session, org_id, user_id, title="Restricted vulnerability assessment",
        disclosure="restricted", as_policy=False,
    )
    review = await _review(client, db_session)
    await db_session.commit()

    preview = (
        await client.post(
            f"/api/v1/compliance/security-reviews/{review.id}/export/preview",
            json={
                "output_type": "evidence_pack",
                "include_documents": [str(restricted.id)],
            },
        )
    ).json()
    row = next(
        document
        for document in preview["documents"]
        if document["file_id"] == str(restricted.id)
    )
    assert row["included"] is False

    # And generation refuses even if a client ignores the preview.
    generated = await client.post(
        f"/api/v1/compliance/security-reviews/{review.id}/export",
        json={
            "output_type": "evidence_pack",
            "include_documents": [str(restricted.id)],
            "confirm_digest": preview["digest"],
        },
    )
    # Refused, or generated without it — never generated with it.
    if generated.status_code == 201:
        export = await db_session.get(ReviewExport, uuid.UUID(generated.json()["id"]))
        titles = [d["title"] for d in export.contents.get("documents", [])]
        assert "Restricted vulnerability assessment" not in titles


@pytest.mark.anyio
async def test_a_superseded_version_is_not_offered_as_current(
    client: AsyncClient, db_session: AsyncSession
) -> None:
    """§12: sending v1.0 as the current policy is worse than sending nothing."""
    org_id, user_id = await _signup(client, f"ver-{uuid.uuid4().hex[:8]}@example.com", "Ver Org")
    old = await _document(
        db_session, org_id, user_id, title="Information Security Policy v1",
        disclosure="customer_shareable", version="1.0",
    )
    new = await _document(
        db_session, org_id, user_id, title="Information Security Policy v2",
        disclosure="customer_shareable", version="2.0",
    )
    old.superseded_by_id = new.id
    new.supersedes_id = old.id
    review = await _review(client, db_session)
    await db_session.commit()

    preview = (
        await client.post(
            f"/api/v1/compliance/security-reviews/{review.id}/export/preview",
            json={"output_type": "evidence_pack"},
        )
    ).json()
    by_title = {document["title"]: document for document in preview["documents"]}

    assert by_title["Information Security Policy v1"]["eligible"] is False
    assert "Superseded" in by_title["Information Security Policy v1"]["reason"]
    assert by_title["Information Security Policy v2"]["included"] is True


# --- §28/§19/§20/§21 generation, snapshot and history ------------------


@pytest.mark.anyio
async def test_generation_requires_confirming_what_was_shown(
    client: AsyncClient, db_session: AsyncSession
) -> None:
    """Between reading a preview and clicking Generate, somebody else can
    approve an answer. The digest is what makes this "send what I saw"."""
    _org, _user = await _signup(client, f"conf-{uuid.uuid4().hex[:8]}@example.com", "Conf Org")
    review = await _review(client, db_session)

    stale = await client.post(
        f"/api/v1/compliance/security-reviews/{review.id}/export",
        json={"output_type": "summary", "confirm_digest": "a" * 64},
    )
    assert stale.status_code == 422
    assert "changed since this preview" in stale.json()["error"]["message"]


@pytest.mark.anyio
async def test_generating_freezes_the_contents_and_records_it(
    client: AsyncClient, db_session: AsyncSession
) -> None:
    """§19/§23/§28: an immutable record, and an audit event naming what went."""
    _org, _user = await _signup(client, f"gen-{uuid.uuid4().hex[:8]}@example.com", "Gen Org")
    review = await _review(client, db_session)

    preview = (
        await client.post(
            f"/api/v1/compliance/security-reviews/{review.id}/export/preview",
            json={"output_type": "summary"},
        )
    ).json()
    generated = await client.post(
        f"/api/v1/compliance/security-reviews/{review.id}/export",
        json={"output_type": "summary", "confirm_digest": preview["digest"]},
    )
    assert generated.status_code == 201, generated.text
    body = generated.json()
    assert body["version"] == 1
    assert body["file_format"] == "pdf"
    assert body["answer_count"] == 3
    assert body["file_id"] is not None
    # §27: generated is not shared.
    assert body["shared_at"] is None

    export = await db_session.get(ReviewExport, uuid.UUID(body["id"]))
    await db_session.refresh(export)
    assert len(export.contents["answers"]) == 3
    assert export.contents["disclosures"], "the qualification was carried into the snapshot"
    assert export.source_digest == preview["digest"]

    event = (
        await db_session.execute(
            select(AuditEvent).where(
                AuditEvent.action == "compliance.review_export.summary_generated"
            )
        )
    ).scalars().one()
    assert event.actor_user_id is not None
    assert event.after_state["customer"] == "Agora Origin"
    assert event.after_state["version"] == 1
    assert event.after_state["answers"] == 3


@pytest.mark.anyio
async def test_regenerating_adds_a_version_and_keeps_the_old_one(
    client: AsyncClient, db_session: AsyncSession
) -> None:
    """§21/§31: the customer has their copy. Ours must still match it."""
    _org, _user = await _signup(client, f"regen-{uuid.uuid4().hex[:8]}@example.com", "Regen Org")
    review = await _review(client, db_session)

    async def generate() -> dict:
        preview = (
            await client.post(
                f"/api/v1/compliance/security-reviews/{review.id}/export/preview",
                json={"output_type": "summary"},
            )
        ).json()
        response = await client.post(
            f"/api/v1/compliance/security-reviews/{review.id}/export",
            json={"output_type": "summary", "confirm_digest": preview["digest"]},
        )
        assert response.status_code == 201, response.text
        return response.json()

    first = await generate()
    first_contents = (
        await db_session.get(ReviewExport, uuid.UUID(first["id"]))
    ).contents["answers"][0]["answer"]

    # An answer is reworded and re-approved.
    questionnaire = (
        await db_session.execute(select(Questionnaire).limit(1))
    ).scalars().one()
    question = (
        await db_session.execute(
            select(QuestionnaireQuestion)
            .where(QuestionnaireQuestion.questionnaire_id == questionnaire.id)
            .order_by(QuestionnaireQuestion.order_index)
            .limit(1)
        )
    ).scalars().one()
    await client.patch(
        f"/api/v1/compliance/questionnaires/{questionnaire.id}/questions/"
        f"{question.id}/answer",
        json={
            "answer_text": "Yes, and the tabletop exercise was completed in August.",
            "status": "approved",
        },
    )

    history = (
        await client.get(f"/api/v1/compliance/security-reviews/{review.id}/exports")
    ).json()
    assert history[0]["stale"] is True, "the screen should say the review moved"

    second = await generate()
    assert second["version"] == 2

    history = (
        await client.get(f"/api/v1/compliance/security-reviews/{review.id}/exports")
    ).json()
    assert len(history) == 2
    assert {row["version"] for row in history} == {1, 2}

    # Version 1 is untouched: it still says what was sent.
    unchanged = (
        await db_session.get(ReviewExport, uuid.UUID(first["id"]))
    ).contents["answers"][0]["answer"]
    assert unchanged == first_contents
    assert "tabletop exercise was completed" not in unchanged


@pytest.mark.anyio
async def test_generating_is_not_sharing(
    client: AsyncClient, db_session: AsyncSession
) -> None:
    """§27: a file existing is not a delivery somebody performed."""
    _org, _user = await _signup(client, f"share-{uuid.uuid4().hex[:8]}@example.com", "Share Org")
    review = await _review(client, db_session)
    preview = (
        await client.post(
            f"/api/v1/compliance/security-reviews/{review.id}/export/preview",
            json={"output_type": "summary"},
        )
    ).json()
    export = (
        await client.post(
            f"/api/v1/compliance/security-reviews/{review.id}/export",
            json={"output_type": "summary", "confirm_digest": preview["digest"]},
        )
    ).json()
    assert export["shared_at"] is None

    marked = await client.post(
        f"/api/v1/compliance/security-reviews/{review.id}/exports/{export['id']}/shared"
    )
    assert marked.status_code == 200
    assert marked.json()["shared_at"] is not None

    # A retry is not a second delivery.
    again = await client.post(
        f"/api/v1/compliance/security-reviews/{review.id}/exports/{export['id']}/shared"
    )
    # Not compared as strings: the same instant serialises with and
    # without a trailing Z depending on whether the row was just written
    # or read back. What matters is that no second delivery was recorded,
    # which the event count below asserts.
    assert again.json()["shared_at"] is not None
    events = list(
        (
            await db_session.execute(
                select(AuditEvent).where(
                    AuditEvent.action == "compliance.review_export.shared"
                )
            )
        )
        .scalars()
        .all()
    )
    assert len(events) == 1


# --- §17 the files are real --------------------------------------------


@pytest.mark.anyio
async def test_each_output_produces_a_real_openable_file(
    client: AsyncClient, db_session: AsyncSession
) -> None:
    """§17: no format is claimed until it opens."""
    org_id, user_id = await _signup(client, f"file-{uuid.uuid4().hex[:8]}@example.com", "File Org")
    await _document(
        db_session, org_id, user_id, title="Information Security Policy",
        disclosure="customer_shareable",
    )
    review = await _review(client, db_session)
    await db_session.commit()

    async def preview_for(output_type: str) -> dict:
        return (
            await client.post(
                f"/api/v1/compliance/security-reviews/{review.id}/export/preview",
                json={"output_type": output_type},
            )
        ).json()

    built = await review_export.build_preview(
        db_session, review=review, output_type="summary"
    )

    pdf = render_summary_pdf(built)
    assert pdf[:5] == b"%PDF-"
    text = "\n".join(page.extract_text() or "" for page in PdfReader(BytesIO(pdf)).pages)
    # §14: the qualification survives into the document, verbatim.
    assert "tabletop exercise" in text
    assert "however" in text
    # §8: no internal identifiers reach the customer.
    assert not any(len(w) == 36 and w.count("-") == 4 for w in text.split())
    assert "confidence" not in text.lower()

    xlsx = render_questionnaire_xlsx(built)
    workbook = load_workbook(BytesIO(xlsx))
    assert workbook.sheetnames == ["Questionnaire", "About"]
    sheet = workbook["Questionnaire"]
    assert sheet.max_row == 4  # header plus three answers
    assert sheet.cell(row=1, column=1).value == "Question"

    index = render_evidence_index(built)
    assert "EVIDENCE INDEX" in index
    assert "Information Security Policy" in index
    assert "Restricted" not in index


@pytest.mark.anyio
async def test_the_evidence_pack_contains_exactly_what_was_previewed(
    client: AsyncClient, db_session: AsyncSession
) -> None:
    """§30: the package is the preview, made into a file."""
    org_id, user_id = await _signup(client, f"pack-{uuid.uuid4().hex[:8]}@example.com", "Pack Org")
    await _document(
        db_session, org_id, user_id, title="Information Security Policy",
        disclosure="customer_shareable",
    )
    await _document(
        db_session, org_id, user_id, title="Restricted vulnerability assessment",
        disclosure="restricted", as_policy=False,
    )
    review = await _review(client, db_session)
    await db_session.commit()

    preview = (
        await client.post(
            f"/api/v1/compliance/security-reviews/{review.id}/export/preview",
            json={"output_type": "evidence_pack"},
        )
    ).json()
    included = {d["title"] for d in preview["documents"] if d["included"]}
    assert included == {"Information Security Policy"}

    generated = await client.post(
        f"/api/v1/compliance/security-reviews/{review.id}/export",
        json={"output_type": "evidence_pack", "confirm_digest": preview["digest"]},
    )
    assert generated.status_code == 201, generated.text
    assert generated.json()["file_format"] == "zip"
    assert generated.json()["document_count"] == 1

    export = await db_session.get(ReviewExport, uuid.UUID(generated.json()["id"]))
    titles = {document["title"] for document in export.contents["documents"]}
    assert titles == included
    assert "Restricted vulnerability assessment" not in titles


# --- §24 permissions ----------------------------------------------------


@pytest.mark.anyio
async def test_disclosure_changes_are_recorded(
    client: AsyncClient, db_session: AsyncSession
) -> None:
    """"Who marked the architecture diagram shareable, and when" is a
    question somebody asks after it has gone out."""
    org_id, user_id = await _signup(client, f"disc-{uuid.uuid4().hex[:8]}@example.com", "Disc Org")
    document = await _document(
        db_session, org_id, user_id, title="Internal architecture diagram",
        disclosure="internal", as_policy=False,
    )
    await db_session.commit()

    response = await client.patch(
        f"/api/v1/compliance/documents/{document.id}/disclosure",
        json={"disclosure": "customer_shareable"},
    )
    assert response.status_code == 200, response.text
    assert response.json()["changed"] is True

    # A retry is not a second decision.
    again = await client.patch(
        f"/api/v1/compliance/documents/{document.id}/disclosure",
        json={"disclosure": "customer_shareable"},
    )
    assert again.json()["changed"] is False

    events = list(
        (
            await db_session.execute(
                select(AuditEvent).where(
                    AuditEvent.action == "compliance.document.disclosure_changed"
                )
            )
        )
        .scalars()
        .all()
    )
    assert len(events) == 1
    assert events[0].before_state["disclosure"] == "internal"
    assert events[0].after_state["disclosure"] == "customer_shareable"
    assert events[0].actor_user_id is not None
